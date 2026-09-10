import json
import re
from pathlib import Path
from collections import Counter

MAX_CHARS = 120_000
SUPPORTED_EXTENSIONS = {".txt", ".md", ".json", ".csv", ".xlsx", ".xls", ".pdf", ".docx", ".pptx"}


def _clean(text):
    text = re.sub(r"\s+", " ", str(text)).strip()
    return text


def extract_text(path):
    path = Path(path)
    ext = path.suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {ext or 'unknown'}")

    if ext in {".txt", ".md"}:
        return path.read_text(encoding="utf-8", errors="ignore")[:MAX_CHARS]

    if ext == ".json":
        data = json.loads(path.read_text(encoding="utf-8", errors="ignore"))
        return json.dumps(data, ensure_ascii=False, indent=2)[:MAX_CHARS]

    if ext in {".csv", ".xlsx", ".xls"}:
        import pandas as pd
        df = pd.read_csv(path) if ext == ".csv" else pd.read_excel(path)
        return df.to_string(index=False)[:MAX_CHARS]

    if ext == ".pdf":
        import fitz
        with fitz.open(path) as doc:
            text = "\n".join(page.get_text("text") for page in doc)
        return text[:MAX_CHARS]

    if ext == ".docx":
        from docx import Document
        doc = Document(path)
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts)[:MAX_CHARS]

    if ext == ".pptx":
        from pptx import Presentation
        prs = Presentation(path)
        parts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    parts.append(shape.text)
        return "\n".join(parts)[:MAX_CHARS]

    raise ValueError("Unsupported file type")


def split_sentences(text):
    text = text.replace("\n", " ")
    return [s.strip() for s in re.split(r"(?<=[.!?])\s+", text) if len(s.strip()) > 30]


def summarize_text(text, max_sentences=6):
    sentences = split_sentences(text)
    if not sentences:
        words = text.split()
        if not words:
            return "No readable text was found in this file."
        return " ".join(words[:120]) + ("..." if len(words) > 120 else "")

    stop = {"the","and","that","this","with","from","have","were","which","their","there","about","into","your","also","will","would","could","should","for","are","was","has","had","not","but","can","its","our","you","they","them","then","than","these","those","a","an","of","to","in","on","is","it","as","or","by","be","at","we","he","she"}
    words = re.findall(r"[a-zA-Z]{3,}", text.lower())
    freq = Counter(w for w in words if w not in stop)
    if not freq:
        return "\n\n".join(sentences[:max_sentences])

    scored = []
    for idx, sentence in enumerate(sentences):
        sw = re.findall(r"[a-zA-Z]{3,}", sentence.lower())
        score = sum(freq[w] for w in sw if w in freq) / max(1, len(sw))
        scored.append((score, idx, sentence))

    selected = sorted(scored, reverse=True)[:max_sentences]
    selected.sort(key=lambda x: x[1])
    return " ".join(s[2] for s in selected)


def answer_from_file(question, text):
    sentences = split_sentences(text)
    if not sentences:
        return "I couldn't find enough readable sentences in the uploaded file."
    qwords = set(re.findall(r"[a-zA-Z]{3,}", question.lower()))
    scored = []
    for sentence in sentences:
        words = set(re.findall(r"[a-zA-Z]{3,}", sentence.lower()))
        overlap = len(qwords & words)
        if overlap:
            scored.append((overlap, sentence))
    if not scored:
        return "I couldn't find a relevant passage in the uploaded file for that question."
    scored.sort(key=lambda x: x[0], reverse=True)
    return " ".join(s for _, s in scored[:3])
